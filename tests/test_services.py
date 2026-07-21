"""
Unit tests for FileParserService and API routes in SynkAI Sprint 2.
"""

import os
import tempfile
import pytest
from docx import Document
from pypdf import PdfWriter
from fastapi.testclient import TestClient

from backend.main import app
from backend.services.file_parser import FileParserService

client = TestClient(app)


def test_parse_txt_file():
    """Test text parsing from plain TXT file."""
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as f:
        f.write("This is a sample meeting transcript line 1.\nThis is line 2.")
        temp_path = f.name

    try:
        parsed_text = FileParserService.parse_file(temp_path)
        assert "sample meeting transcript line 1" in parsed_text
        assert "This is line 2." in parsed_text
    finally:
        os.remove(temp_path)


def test_parse_docx_file():
    """Test text parsing from Word DOCX file."""
    doc = Document()
    doc.add_paragraph("DOCX Meeting Notes Heading")
    doc.add_paragraph("Item 1: Discuss quarterly roadmap")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        temp_path = f.name

    try:
        doc.save(temp_path)
        parsed_text = FileParserService.parse_file(temp_path)
        assert "DOCX Meeting Notes Heading" in parsed_text
        assert "Discuss quarterly roadmap" in parsed_text
    finally:
        os.remove(temp_path)


def test_upload_route_txt():
    """Test POST /upload endpoint with text file stream."""
    file_content = b"Meeting Transcript: Discussed Sprint 2 scope and API design."
    files = {"file": ("test_transcript.txt", file_content, "text/plain")}

    response = client.post("/upload", files=files)
    assert response.status_code == 201
    json_data = response.json()
    assert json_data["filename"] == "test_transcript.txt"
    assert json_data["character_count"] == len(file_content)
    assert "test_transcript.txt" in json_data["file_path"]


def test_upload_route_invalid_extension():
    """Test POST /upload rejects unsupported file extensions."""
    files = {"file": ("script.exe", b"binary content", "application/octet-stream")}
    response = client.post("/upload", files=files)
    assert response.status_code == 400
    assert "Unsupported file format" in response.json()["detail"]
